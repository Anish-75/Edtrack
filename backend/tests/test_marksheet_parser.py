from openpyxl import Workbook

from marksheet_parser import parse_marksheet


def test_parse_marksheet_returns_student_and_course(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.cell(3, 3, "PR. NO.")
    worksheet.cell(3, 5, "NAME OF THE CANDIDATE")
    worksheet.cell(3, 7, "EXNO")
    worksheet.cell(3, 8, "CMP-100")
    worksheet.cell(3, 13, "SEMESTER")
    worksheet.cell(3, 16, "RESULT")
    worksheet.cell(4, 8, "Programming")

    for column, value in zip(range(8, 12), ("CA", "CE", "GP", "GRADE")):
        worksheet.cell(5, column, value)
    worksheet.cell(5, 13, "TCE")
    worksheet.cell(5, 14, "SGPA")
    worksheet.cell(5, 15, "P/F")

    worksheet.cell(6, 3, 12345)
    worksheet.cell(6, 5, "Test Student")
    worksheet.cell(6, 7, 54321)
    for column, value in zip(range(8, 12), (10, 10, 4, "A")):
        worksheet.cell(6, column, value)
    worksheet.cell(6, 13, 10)
    worksheet.cell(6, 14, 8.4)
    worksheet.cell(6, 15, "PASS")
    worksheet.cell(6, 16, "PASS")

    input_path = tmp_path / "synthetic_marksheet.xlsx"
    workbook.save(input_path)

    students = parse_marksheet(str(input_path))

    assert len(students) == 1
    student = students[0]
    assert student["pr_number"] == "12345"
    assert student["name"] == "Test Student"
    assert student["seat_number"] == "54321"
    assert student["tce"] == 10
    assert student["sgpa"] == 8.4
    assert student["result"] == "PASS"
    assert student["pf"] == "PASS"
    assert student["courses"] == [
        {
            "code": "CMP-100",
            "name": "Programming",
            "ca": 10,
            "ce": 10,
            "gp": 4,
            "grade": "A",
        }
    ]
