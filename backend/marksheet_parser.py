"""
marksheet_parser.py
====================
Fully dynamic parser for Goa University wide-format marksheet Excel files.

Instead of hardcoding column indices, this parser:
  1. Auto-detects the header row by searching for student-info anchor columns.
  2. Scans that row to discover all course codes and summary block positions.
  3. For each course, reads the sub-header row to locate CA / CE / GP / GRADE.
  4. Detects summary columns (TCE, SGPA, Result) by keyword search.

This means the parser works even if:
  - Subjects change between semesters.
  - Columns are added, removed, or re-ordered.
  - A course block spans a different number of columns.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


# ── Constants ─────────────────────────────────────────────────────────────────

_STUDENT_INFO_HEADERS = {
    'SR. NO.', 'SR.NO.', 'RITNO', 'PR. NO.', 'PR.NO.',
    'ABC ID', 'NAME OF THE CANDIDATE', 'M / F', 'M/F', 'EXNO', 'EXAM NO.',
}

_SUMMARY_KEYWORDS = {'SEM', 'SEMESTER', 'CGPA', 'RESULT', 'TCE', 'TOTAL', 'SGPA', 'SEM - I'}

_WANT = {'CA', 'CE', 'GP', 'GRADE'}

_THIN   = Side(style='thin')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_GRADE_FILLS = {
    'O':  PatternFill('solid', fgColor='FFF3CD'),
    'A+': PatternFill('solid', fgColor='D4EDDA'),
    'A':  PatternFill('solid', fgColor='D1ECF1'),
    'B+': PatternFill('solid', fgColor='CCE5FF'),
    'B':  PatternFill('solid', fgColor='E2CFEE'),
    'C':  PatternFill('solid', fgColor='F8D7DA'),
    'D':  PatternFill('solid', fgColor='F8D7DA'),
}

_ANCHOR_COLS = {'PR. NO.', 'PR.NO.', 'NAME OF THE CANDIDATE', 'EXNO'}


def _str(value, *, numeric: bool = False) -> str:
    if pd.isna(value): return ''
    if numeric:
        try: return str(int(float(value)))
        except (ValueError, TypeError): pass
    return str(value).strip()


def _float(value, default: float = 0.0) -> float:
    if pd.isna(value): return default
    try: return float(value)
    except (ValueError, TypeError): return default


def _discover_layout(df: pd.DataFrame) -> dict:
    code_row = None
    for r in range(min(20, len(df))):
        vals = {str(v).strip().upper() for v in df.iloc[r] if not pd.isna(v)}
        if len(vals & _ANCHOR_COLS) >= 2:
            code_row = r; break

    if code_row is None:
        raise ValueError("Could not locate the header row. Expected 'PR. NO.' and 'Name of the Candidate'.")

    name_row = code_row + 1
    sub_row  = code_row + 2

    pr_col = name_col = seat_col = None
    for c, val in enumerate(df.iloc[code_row]):
        v = str(val).strip().upper()
        if v in {'PR. NO.', 'PR.NO.'}: pr_col = c
        elif v == 'NAME OF THE CANDIDATE': name_col = c
        elif v in {'EXNO', 'EXAM NO.'}: seat_col = c

    if pr_col is None: raise ValueError("Could not find 'PR. NO.' column.")
    if name_col is None: raise ValueError("Could not find 'Name of the Candidate' column.")

    data_start = None
    for r in range(code_row + 2, min(code_row + 20, len(df))):
        try: int(float(df.iloc[r, pr_col])); data_start = r; break
        except (ValueError, TypeError): pass

    if data_start is None:
        raise ValueError("Could not find the first student data row.")

    summary_start = None
    course_cols: dict[int, str] = {}

    for c, val in enumerate(df.iloc[code_row]):
        if pd.isna(val): continue
        v = str(val).strip()
        if v.upper() in _STUDENT_INFO_HEADERS: continue
        if any(kw in v.upper() for kw in _SUMMARY_KEYWORDS):
            if summary_start is None: summary_start = c
            continue
        course_cols[c] = v

    sorted_starts = sorted(course_cols.keys())
    courses = []

    for i, start in enumerate(sorted_starts):
        end  = sorted_starts[i + 1] if i + 1 < len(sorted_starts) else (summary_start or df.shape[1])
        code = course_cols[start]
        name = _str(df.iloc[name_row, start])
        if not name or name == code: name = code

        sub_map: dict[str, int] = {}
        for c in range(start, end):
            sv = str(df.iloc[sub_row, c]).strip().upper()
            if sv in _WANT and sv not in sub_map: sub_map[sv] = c

        if not _WANT.issubset(sub_map): continue
        courses.append({'code': code, 'name': name,
                        'ca': sub_map['CA'], 'ce': sub_map['CE'],
                        'gp': sub_map['GP'], 'grade': sub_map['GRADE']})

    if not courses: raise ValueError("No course columns found in the header row.")

    tce_col = sgpa_col = result_col = pf_col = None
    for c in range(summary_start or 0, df.shape[1]):
        sv = str(df.iloc[sub_row,  c]).strip().upper()
        v3 = str(df.iloc[code_row, c]).strip().upper()
        if 'TCE'    in sv and tce_col    is None: tce_col    = c
        if 'SGPA'   in sv and sgpa_col   is None: sgpa_col   = c
        if 'P/F'    in sv and pf_col     is None: pf_col     = c
        if 'RESULT' in v3 and result_col is None: result_col = c

    return dict(code_row=code_row, name_row=name_row, sub_row=sub_row,
                data_start=data_start, pr_col=pr_col, name_col=name_col,
                seat_col=seat_col, courses=courses, tce_col=tce_col,
                sgpa_col=sgpa_col, result_col=result_col, pf_col=pf_col)


def validate_marksheet_structure(input_file_path: str) -> None:
    try:
        df = pd.read_excel(input_file_path, header=None)
    except Exception as exc:
        raise ValueError(f"Could not open file as Excel: {exc}") from exc
    _discover_layout(df)


def parse_marksheet(input_file_path: str) -> list[dict]:
    try:
        df     = pd.read_excel(input_file_path, header=None)
        layout = _discover_layout(df)

        pr_col=layout['pr_col']; name_col=layout['name_col']; seat_col=layout['seat_col']
        data_start=layout['data_start']; courses=layout['courses']
        tce_col=layout['tce_col']; sgpa_col=layout['sgpa_col']
        result_col=layout['result_col']; pf_col=layout['pf_col']

        students: list[dict] = []
        for r in range(data_start, len(df)):
            row = df.iloc[r]
            if pd.isna(row[pr_col]) or _str(row[pr_col]) == '': break
            s = {
                'pr_number':   _str(row[pr_col], numeric=True),
                'name':        _str(row[name_col]),
                'seat_number': _str(row[seat_col]) if seat_col is not None else '',
                'courses':     [],
                'tce':         int(_float(row[tce_col]))       if tce_col    is not None else 0,
                'sgpa':        round(_float(row[sgpa_col]), 2) if sgpa_col   is not None else 0.0,
                'result':      _str(row[result_col])           if result_col is not None else '',
                'pf':          _str(row[pf_col])               if pf_col     is not None else '',
            }
            for cs in courses:
                s['courses'].append({
                    'code':  cs['code'], 'name': cs['name'],
                    'ca':    int(_float(row[cs['ca']])), 'ce': int(_float(row[cs['ce']])),
                    'gp':    int(_float(row[cs['gp']])), 'grade': _str(row[cs['grade']]),
                })
            students.append(s)
        return students

    except ValueError: raise
    except Exception as exc:
        raise Exception(f"Error parsing marksheet: {exc}") from exc


def create_individual_report(student: dict, output_path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    bold = Font(size=11, bold=True)
    normal = Font(size=10)
    center = Alignment(horizontal='center', vertical='center')

    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Student Info (MERGED STYLE) ───────
    def label_value(row, label, value):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=11)

        ws.cell(row=row, column=1, value=label).font = normal
        ws.cell(row=row, column=3, value=value).font = bold

    label_value(3, "Name of the Candidate", student["name"])
    label_value(4, "SEAT NUMBER", student["seat_number"])
    label_value(5, "PR NUMBER", student["pr_number"])
    label_value(6, "DEPARTMENT", "ALL DEPARTMENT")
    label_value(7, "COLLEGE", "SHREE RAYESHWAR INSTITUTE OF ENGINEERING & INFORMATION TECHNOLOGY")
    label_value(8, "SCHEME", "RC 2024-25")
    # ── Table Header (MERGED) ─────────────
    start_row = 10

    headers = [
        ("A", "A", "Course Code"),
        ("B", "F", "Nomenclature of the Course"),
        ("G", "G", "Credits Assigned"),
        ("H", "H", "Credits Earned"),
        ("I", "I", "Grade Point"),
        ("J", "J", "Letter Grade"),
    ]

    for col_start, col_end, text in headers:
        ws.merge_cells(f"{col_start}{start_row}:{col_end}{start_row}")
        cell = ws[f"{col_start}{start_row}"]
        cell.value = text
        cell.font = bold
        cell.alignment = center

        # Apply border to merged region
        for col in range(ord(col_start), ord(col_end) + 1):
            ws.cell(row=start_row, column=col - 64).border = border

    # ── Course Rows ───────────────────────
    row = start_row + 1

    for course in student["courses"]:
        ws.merge_cells(f"A{row}:A{row}")
        ws.merge_cells(f"B{row}:F{row}")

        ws[f"A{row}"] = course["code"]
        ws[f"B{row}"] = course["name"]
        ws[f"G{row}"] = course["ca"]
        ws[f"H{row}"] = course["ce"]
        ws[f"I{row}"] = course["gp"]
        ws[f"J{row}"] = course["grade"]

        # Apply borders
        for col in range(1, 11):
            ws.cell(row=row, column=col).border = border

        row += 1

    # ── Total Row ─────────────────────────
    row += 1
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"] = "Total Credits Earned"
    ws[f"A{row}"].font = bold
    ws[f"G{row}"] = student["tce"]

    for col in range(1, 11):
        ws.cell(row=row, column=col).border = border

    # ── Column Widths ─────────────────────
    widths = {
        "A": 10.56, "B": 7.56,
        "C": 7.33, "D": 7.33, "E": 9.78, "F": 5.89,
        "G": 7.89, "H": 7.89, "I": 7.89, "J": 7.89,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width
        
    ws.row_dimensions[10].height = 26.3
    
    for r in range(11, 20):   
        ws.row_dimensions[r].height = 21.8

    wb.save(output_path)
    return output_path